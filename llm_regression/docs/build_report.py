"""生成汇报文档（DOCX）。

用法：
    .venv\\Scripts\\python.exe llm_regression/docs/build_report.py
输出：
    llm_regression/docs/AI大模型对话接口自动化回归测试_项目报告.docx

设计说明：
- 中文字体设为「微软雅黑」（含东亚字体属性，避免打开后变宋体/方框）；
- 标题分级 + 表格化呈现，方便直接打印或转 PDF；
- 所有数字均来自真实产物（见文档中的「数据来源」说明）。
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "微软雅黑"
DARK = RGBColor(0x1F, 0x2A, 0x37)
GRAY = RGBColor(0x5B, 0x66, 0x73)
ACCENT = RGBColor(0x1F, 0x6F, 0xEB)
RED = RGBColor(0xC0, 0x1B, 0x1B)

OUT = Path(__file__).resolve().parent / "AI大模型对话接口自动化回归测试_项目报告.docx"


# ------------------------------------------------------------------ 基础工具
def set_font(run, size=10.5, bold=False, color=DARK, name=FONT) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    # 中文字体必须单独设置 eastAsia，否则 Word 会用默认宋体
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def add_para(doc, text="", *, size=10.5, bold=False, color=DARK, align=None,
             space_before=0, space_after=6, indent_first=False, line_spacing=1.4):
    para = doc.add_paragraph()
    if align is not None:
        para.alignment = align
    fmt = para.paragraph_format
    fmt.space_before = Pt(space_before)
    fmt.space_after = Pt(space_after)
    fmt.line_spacing = line_spacing
    if indent_first:
        fmt.first_line_indent = Pt(size * 2)
    if text:
        set_font(para.add_run(text), size=size, bold=bold, color=color)
    return para


def add_heading(doc, text, level=1):
    """自定义标题：比默认样式更紧凑，且字体受控。"""
    sizes = {1: 16, 2: 13.5, 3: 11.5}
    before = {1: 16, 2: 12, 3: 8}
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(before[level])
    para.paragraph_format.space_after = Pt(6)
    set_font(para.add_run(text), size=sizes[level], bold=True,
             color=ACCENT if level <= 2 else DARK)
    return para


def add_bullet(doc, text, *, size=10.5, bold=False, level=0):
    para = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    para.paragraph_format.space_after = Pt(3)
    para.paragraph_format.line_spacing = 1.35
    set_font(para.add_run(text), size=size, bold=bold)
    return para


def add_table(doc, headers, rows, *, widths=None, header_fill="1F6FEB", zebra=True):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    hdr = table.rows[0].cells
    for index, title in enumerate(headers):
        hdr[index].text = ""
        para = hdr[index].paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_after = Pt(2)
        set_font(para.add_run(title), size=10, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        shade = hdr[index]._tc.get_or_add_tcPr()
        el = shade.makeelement(qn("w:shd"), {qn("w:fill"): header_fill})
        shade.append(el)

    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for col_index, value in enumerate(row):
            cells[col_index].text = ""
            para = cells[col_index].paragraphs[0]
            para.paragraph_format.space_after = Pt(2)
            para.paragraph_format.line_spacing = 1.25
            set_font(para.add_run(str(value)), size=9.5)
            if zebra and row_index % 2 == 1:
                shade = cells[col_index]._tc.get_or_add_tcPr()
                el = shade.makeelement(qn("w:shd"), {qn("w:fill"): "F5F8FC"})
                shade.append(el)

    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Cm(width)
    add_para(doc, "", space_after=4)
    return table


def add_caption(doc, text):
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_after = Pt(10)
    set_font(para.add_run(text), size=9, color=GRAY)


# ------------------------------------------------------------------ 文档内容
def build() -> None:
    doc = Document()

    # 页面设置
    section = doc.sections[0]
    section.top_margin = Cm(2.4)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.6)
    section.right_margin = Cm(2.6)

    # 正文默认字体
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    # ============================================================ 封面
    add_para(doc, "", space_after=60)
    add_para(doc, "AI 大模型对话接口", size=26, bold=True, color=DARK,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    add_para(doc, "自动化回归测试", size=26, bold=True, color=ACCENT,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
    add_para(doc, "项目报告", size=15, color=GRAY,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=50)
    add_para(doc, "从「人工逐条验证」到「可重复、可度量、可追溯」",
             size=11.5, color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=70)

    add_table(
        doc,
        ["项目", "内容"],
        [
            ["汇报人", ""],
            ["汇报日期", ""],
            ["被测对象", "AI 大模型对话接口（OpenAI 兼容，HTTPS）"],
            ["交付形式", "可运行的自动化回归框架 + 测试报告 + 缺陷单"],
            ["代码位置", "llm_regression/（分支 test/llm-regression-framework）"],
        ],
        widths=[3.2, 11.2],
    )

    doc.add_page_break()

    # ============================================================ 一、项目概述
    add_heading(doc, "一、项目概述", 1)

    add_heading(doc, "1.1 背景", 2)
    add_para(doc,
             "被测对象是一个 AI 大模型对话接口，该接口版本迭代频繁、入参组合多，"
             "且返回内容由模型实时生成。原有的回归方式是测试人员手工在 Postman 中逐条构造请求、"
             "肉眼比对返回结果，单次回归耗时约 30 分钟，且存在三个突出问题：", indent_first=True)
    add_bullet(doc, "覆盖依赖个人经验：空输入、超长文本、非法字符、异常编码等边界与异常场景，"
                    "是否覆盖、覆盖到什么程度完全取决于测试人员经验，容易漏测；")
    add_bullet(doc, "过程无留痕：回归过程没有记录，出现问题时无法说明当时的测试数据与预期依据；")
    add_bullet(doc, "结果不可度量：模型返回内容每次不同，人工判断标准不统一，"
                    "回归结论难以作为发布依据。")

    add_heading(doc, "1.2 项目目标", 2)
    add_para(doc, "把「人工点一遍」升级为「可重复、可度量、可追溯」的自动化回归，具体拆解为四项：",
             indent_first=True)
    add_table(
        doc,
        ["目标", "衡量方式"],
        [
            ["可重复", "一条命令即可执行全部场景，结果稳定不依赖人工判断"],
            ["可度量", "除状态码外，校验响应结构、字段类型、业务错误码与响应耗时"],
            ["可追溯", "每次执行自动生成报告归档，失败用例自动留存请求与响应证据"],
            ["可扩展", "新增接口或新增场景只需补充配置与数据，不需要改动代码"],
        ],
        widths=[2.8, 11.6],
    )

    add_heading(doc, "1.3 交付成果一览", 2)
    add_table(
        doc,
        ["交付物", "数量 / 结果", "说明"],
        [
            ["自动化回归框架", "1 套（4 层结构）", "配置层 / 请求层 / 数据层 / 断言层"],
            ["回归用例", "Mock 30 条 / 真实 18 条", "四类场景，全部通过"],
            ["断言规则", "8 类", "含耗时阈值，把性能回退纳入判定"],
            ["可视化报告", "每次执行自动生成", "HTML + JUnit XML，失败自动留证"],
            ["缺陷单", "6 条", "已提交缺陷管理系统并附实测证据"],
            ["自动化流水线", "2 个环节", "Mock 门禁 + 真实接口回归（可选触发）"],
        ],
        widths=[3.4, 4.4, 6.6],
    )
    add_caption(doc, "表 1-1　交付成果一览（数据来源：项目仓库实际产物）")

    # ============================================================ 二、解决方案
    add_heading(doc, "二、解决方案设计", 1)

    add_heading(doc, "2.1 整体思路", 2)
    add_para(doc,
             "核心思路是「把测试经验固化成数据，而不是留在个人脑子里」。"
             "为此做了两个关键决策：", indent_first=True)
    add_para(doc, "第一，用等价类划分与边界值分析把输入域拆成四类场景，"
                  "每一类都落成可参数化的数据文件；", indent_first=True)
    add_para(doc, "第二，自建一个 Mock 接口作为被测对象的「契约实现」，"
                  "用于稳定复现真实接口无法稳定触发的异常分支（超时、服务端错误、瞬时故障、编码异常）。",
             indent_first=True)

    add_heading(doc, "2.2 四类测试场景", 2)
    add_table(
        doc,
        ["场景", "设计方法", "覆盖要点"],
        [
            ["正常输入", "有效等价类", "中文提问、中英文混合，验证 200 与结构完整"],
            ["空输入", "边界值（0 字符）+ 无效等价类",
             "空串、纯空格、纯换行、字段缺失、字段类型错误"],
            ["超长文本", "边界值三点法", "999 / 1000 / 1001 字符夹住上限，另加 10000 字符极端值"],
            ["非法字符", "无效等价类 + 异常分支",
             "NUL 字符、控制字符、残缺 JSON、非 UTF-8 编码、XSS payload"],
        ],
        widths=[2.4, 4.0, 8.0],
    )
    add_caption(doc, "表 2-1　四类场景与用例设计方法")

    add_heading(doc, "2.3 框架分层结构", 2)
    add_para(doc, "框架按职责分为四层，层与层之间通过配置与数据解耦：", indent_first=True)
    add_table(
        doc,
        ["层级", "职责", "带来的价值"],
        [
            ["配置层", "接口地址、鉴权、超时、重试策略、耗时基线", "更换环境或接口不改代码"],
            ["请求层", "统一封装会话、鉴权头、超时与重试，并记录完整调用留痕",
             "策略集中管理，异常处理一致"],
            ["数据层", "以 YAML 存放用例数据", "新增用例 = 新增一行数据"],
            ["断言层", "8 类可度量断言规则", "失败信息含期望值、实际值与上下文"],
        ],
        widths=[2.2, 6.6, 5.6],
    )
    add_caption(doc, "表 2-2　四层结构职责划分")

    add_heading(doc, "2.4 断言规则（8 类）", 2)
    add_para(doc,
             "断言不只判断「接口通了」，而是判断「结果对不对、快不快」。"
             "这一点是本项目与手工回归的核心差异。", indent_first=True)
    add_table(
        doc,
        ["序号", "断言项", "说明"],
        [
            ["1", "HTTP 状态码", "精确匹配预期状态码"],
            ["2", "响应 JSON 结构", "校验必需字段是否存在"],
            ["3", "关键字段类型", "校验字段类型是否符合契约"],
            ["4", "业务错误码", "校验业务层的错误标识，而非仅看 HTTP 层"],
            ["5", "生成内容非空", "成功响应必须返回可用的生成内容，含最小长度校验"],
            ["6", "响应耗时基线", "把性能回退纳入回归判定：接口通了但慢一倍同样判失败"],
            ["7", "重试次数", "验证重试策略确实生效，而不只是「碰巧成功」"],
            ["8", "超时分支", "验证客户端在超时预算内正确快速失败"],
        ],
        widths=[1.4, 3.6, 9.4],
    )
    add_caption(doc, "表 2-3　八类断言规则")

    add_heading(doc, "2.5 不稳定依赖的治理", 2)
    add_para(doc,
             "AI 接口天然存在三类不稳定：响应慢、依赖网络、异常分支无法按需触发。"
             "本项目通过自建 Mock 服务实现可控：", indent_first=True)
    add_table(
        doc,
        ["不稳定因素", "治理手段"],
        [
            ["服务端慢响应", "Mock 支持按标记注入延迟，用于验证耗时阈值与超时分支"],
            ["服务端持续报错", "Mock 支持持续返回 5xx，验证重试上限"],
            ["瞬时故障", "Mock 支持「前 N 次失败、之后成功」，验证重试自愈能力"],
            ["请求体编码异常", "Mock 支持接收非 UTF-8 字节，验证解码异常的处理路径"],
        ],
        widths=[4.0, 10.4],
    )
    add_caption(doc, "表 2-4　不稳定依赖及其治理手段")

    # ============================================================ 三、执行结果
    add_heading(doc, "三、执行结果", 1)

    add_heading(doc, "3.1 回归结果", 2)
    add_table(
        doc,
        ["执行模式", "用例数", "结果", "耗时", "说明"],
        [
            ["Mock 模式", "30", "全部通过", "约 2 秒", "离线可跑，结果确定"],
            ["真实接口", "18", "全部通过", "约 18 秒", "连接真实大模型服务，含网络波动"],
        ],
        widths=[2.6, 1.8, 2.4, 2.2, 5.4],
    )
    add_caption(doc, "表 3-1　回归执行结果（多次执行结果一致）")

    add_heading(doc, "3.2 效率对比", 2)
    add_table(
        doc,
        ["环节", "人工方式", "自动化方式"],
        [
            ["构造请求", "逐条手工填写参数，含超长文本粘贴", "数据文件参数化，零人工"],
            ["结果比对", "肉眼逐条核对返回", "8 类断言自动判定"],
            ["留痕记录", "人工记录或截图，易遗漏", "自动生成报告，失败自动留存证据"],
            ["汇总整理", "人工整理汇总", "自动输出 HTML 报告与 JUnit XML"],
            ["单次总耗时", "约 30 分钟", "约 2 秒"],
        ],
        widths=[2.6, 5.6, 6.2],
    )
    add_caption(doc, "表 3-2　人工方式与自动化方式对比")

    add_heading(doc, "3.3 自动化流水线", 2)
    add_para(doc,
             "回归已接入持续集成流水线，包含两个环节：", indent_first=True)
    add_table(
        doc,
        ["环节", "触发方式", "作用"],
        [
            ["Mock 回归门禁", "代码提交 / 合并请求自动触发",
             "约 20 秒完成 30 条用例；不通过则拦截代码合入"],
            ["真实接口回归", "手动触发（可选）",
             "对真实服务执行 18 条用例，产物自动归档"],
        ],
        widths=[3.0, 5.0, 6.4],
    )
    add_caption(doc, "表 3-3　流水线环节说明")
    add_para(doc,
             "为验证门禁确实有效，专门做了一次反向验证：故意将一条用例的期望值改错并提交，"
             "流水线如期变红并拦截；改回后恢复通过。该次失败记录保留在流水线历史中，"
             "可作为门禁有效性的证据。", indent_first=True)

    # ============================================================ 四、质量发现
    add_heading(doc, "四、测试发现", 1)
    add_para(doc,
             "本次工作最有价值的产出是通过真实接口回归发现了 6 个问题，"
             "均已提交缺陷管理系统，每条包含前置条件、复现步骤、实际结果、期望结果与实测证据。"
             "以下按严重程度排列：", indent_first=True)
    add_table(
        doc,
        ["编号", "严重程度", "问题描述", "影响"],
        [
            ["1", "高", "存在随机连接中断：相同请求偶发直接断开，重试即可成功",
             "导致自动化回归结果不可信，无法作为发布门禁"],
            ["2", "中", "缺少输入参数校验：空值、纯空白、控制字符均被当作正常请求处理",
             "无效请求进入模型，浪费算力，且掩盖调用方错误"],
            ["3", "中", "缺少长度上限校验：10000 字符输入仍正常处理",
             "存在资源耗尽风险，且无明确的容量契约"],
            ["4", "中", "用户输入被原样写入响应内容，需评估输出编码责任",
             "若由服务端直接渲染为 HTML，存在跨站脚本风险"],
            ["5", "低", "畸形请求只返回纯文本错误，没有结构化错误体",
             "调用方无法用统一逻辑解析错误"],
            ["6", "低", "HTTP 状态码不统一：结构错误 422、业务校验 400、模型名非法 400",
             "调用方无法通过状态码区分错误类别"],
        ],
        widths=[1.2, 1.8, 6.4, 5.0],
    )
    add_caption(doc, "表 4-1　测试发现的 6 个问题")

    add_heading(doc, "4.1 一个需要说明的测试设计问题", 2)
    add_para(doc,
             "在验证输出编码时，最初的断言设计是「响应内容中不得出现可执行脚本片段」。"
             "但实测发现，当输入包含脚本片段时，模型会在回答中复述该内容，"
             "导致该断言必然误报——而模型复述用户输入本身是正常行为。", indent_first=True)
    add_para(doc,
             "据此得出结论：用关键词匹配判定跨站脚本风险在 AI 场景下不可靠。"
             "该用例已降级为「确认服务未报错且返回可用内容」，"
             "真正的输出编码验证需要在客户端渲染层单独设计用例。"
             "这是本项目的一个已知空白，已在第五章说明。", indent_first=True)

    # ============================================================ 五、不足
    add_heading(doc, "五、已知不足与改进方向", 1)
    add_table(
        doc,
        ["不足", "说明", "改进方向"],
        [
            ["安全断言不够严谨",
             "关键词匹配判定脚本注入会因模型复述输入而误报，相关断言已降级",
             "在客户端渲染层单独设计用例，或引入人工评估环节"],
            ["性能基线为经验值",
             "耗时阈值尚未按真实响应时间分布做统计标定",
             "连续采集多轮真实回归耗时，按分位数重新设定阈值"],
            ["权限场景未覆盖",
             "密钥错误等鉴权失败分支尚未纳入回归",
             "补充用例数据即可接入，接口与框架均已支持"],
            ["真实接口回归未纳入强制门禁",
             "当前仅 Mock 回归作为必过门禁，真实回归需手动触发",
             "待连接中断问题治理完成后，再评估纳入强制门禁"],
        ],
        widths=[3.0, 6.0, 5.4],
    )
    add_caption(doc, "表 5-1　已知不足与改进方向")

    # ============================================================ 六、结论
    add_heading(doc, "六、结论", 1)
    add_para(doc,
             "本项目完成了从方案设计到交付的闭环：将人工回归升级为可重复、可度量、可追溯的"
             "自动化回归体系，并已接入持续集成流水线。", indent_first=True)
    add_para(doc, "主要成果：", indent_first=True)
    add_bullet(doc, "效率：单次回归由人工约 30 分钟缩短至约 2 秒，且可反复执行、结果稳定；")
    add_bullet(doc, "覆盖：空输入、超长文本、非法字符三类边界与异常场景纳入常态化回归，"
                    "不再依赖个人经验；")
    add_bullet(doc, "留痕：每次执行自动生成可视化报告并归档，失败用例自动留存请求参数、"
                    "响应内容与耗时，缺陷单可直接附上复现信息；")
    add_bullet(doc, "质量：真实接口回归发现 6 个问题并提交跟踪，其中随机连接中断为高优先级；")
    add_bullet(doc, "可持续：框架分层设计，新增接口只需补充配置与数据；"
                    "Mock 回归已作为代码合入的必过门禁。")
    add_para(doc, "一句话结论：", indent_first=True, space_before=6)
    para = add_para(doc, "自动化回归体系已跑通并接入流水线，用它测出被测接口存在 6 个问题，"
                         "其中「随机连接中断」为高优先级。", size=11.5, bold=True, color=ACCENT)

    norm = doc.add_paragraph()
    norm.paragraph_format.left_indent = Cm(0.8)
    set_font(norm.add_run("需要说明的是，「回归通过」与「被测服务健康」是两件事："
                          "回归通过表示接口行为与已记录的契约一致，"
                          "而接口本身存在的问题已单独作为缺陷提交跟踪。"),
             size=10, color=GRAY)

    # ============================================================ 附录
    doc.add_page_break()
    add_heading(doc, "附录 A　关键数据与复现方式", 1)

    add_heading(doc, "A.1 关键数据", 2)
    add_table(
        doc,
        ["项目", "数值"],
        [
            ["Mock 用例数 / 结果", "30 条 / 全部通过（约 2 秒）"],
            ["真实接口用例数 / 结果", "18 条 / 全部通过（约 18 秒）"],
            ["四类场景", "正常输入、空输入、超长文本、非法字符"],
            ["断言规则", "8 类"],
            ["测试发现问题", "6 条"],
            ["效率提升", "约 30 分钟 → 约 2 秒"],
        ],
        widths=[5.0, 9.4],
    )
    add_caption(doc, "表 A-1　关键数据")

    add_heading(doc, "A.2 复现方式", 2)
    add_para(doc, "在项目根目录执行以下命令即可复现全部结果：", indent_first=True)
    code_style = doc.add_paragraph()
    code_style.paragraph_format.left_indent = Cm(0.6)
    code_style.paragraph_format.space_after = Pt(4)
    set_font(code_style.add_run("# Mock 回归（离线，约 2 秒）"), size=9.5, color=GRAY)
    for line in (".venv\\Scripts\\python.exe llm_regression\\run_tests.py",):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.space_after = Pt(8)
        set_font(p.add_run(line), size=9.5)

    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.6)
    p.paragraph_format.space_after = Pt(4)
    set_font(p.add_run("# 真实接口回归（需外网，约 18 秒）"), size=9.5, color=GRAY)
    for line in (".venv\\Scripts\\python.exe llm_regression\\run_tests.py --env real",):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.space_after = Pt(8)
        set_font(p.add_run(line), size=9.5)

    add_para(doc, "执行完成后，可视化报告位于：", indent_first=True, space_before=6)
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.6)
    set_font(p.add_run("llm_regression\\reports\\latest_mock.html（Mock）"), size=9.5)
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.6)
    set_font(p.add_run("llm_regression\\reports\\latest_real.html（真实接口）"), size=9.5)

    add_heading(doc, "A.3 数据来源说明", 2)
    add_para(doc,
             "本报告中的全部数据均来自项目实际产物，可逐条复核：", indent_first=True)
    add_table(
        doc,
        ["数据", "来源"],
        [
            ["用例数量与通过情况", "reports/junit_mock.xml、reports/junit_real.xml"],
            ["执行耗时", "上述 JUnit 报告中的 time 字段"],
            ["缺陷内容与编号", "缺陷管理系统中的对应记录"],
            ["流水线结果", "代码托管平台的流水线运行历史"],
        ],
        widths=[4.6, 9.8],
    )
    add_caption(doc, "表 A-2　数据来源对照")

    doc.save(str(OUT))
    print(f"已生成：{OUT}")


if __name__ == "__main__":
    build()
